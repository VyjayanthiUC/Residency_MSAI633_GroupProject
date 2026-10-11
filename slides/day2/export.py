#!/usr/bin/env python3
"""Export the deck to .pptx: one full-bleed 1920×1080 screenshot per slide (final state, ?export=1) + speaker notes.

    python3 slides/day2/serve.py &                # preview server (default port 8820)
    uv run --with python-pptx python slides/day2/export.py [out.pptx]

Env: DECK_URL (default http://127.0.0.1:8820/), BROWSER (a Chromium-family binary; auto-detected otherwise).
"""
import html
import os
import shutil
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "FanBot_Day2_Presentation.pptx"
SHOTS = HERE / "shots" / "export"
SHOTS.mkdir(parents=True, exist_ok=True)
URL = os.environ.get("DECK_URL", "http://127.0.0.1:8820/")
BROWSER = os.environ.get("BROWSER") or next((b for b in ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable",
                                                         "brave", "brave-browser", "microsoft-edge", "chrome", "msedge")
                                              if shutil.which(b)), None)
# Windows installs Chrome and Edge outside PATH
WINDOWS = [Path(os.environ.get(v, "")) / p for v in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA")
           for p in ("Google/Chrome/Application/chrome.exe", "Microsoft/Edge/Application/msedge.exe")]
BROWSER = BROWSER or next((str(p) for p in WINDOWS if p.is_file()), None)
if not BROWSER:
    sys.exit("No Chromium-family browser found: set BROWSER=/path/to/chrome")

page = (HERE / "index.html").read_text(encoding="utf-8")
slides = re.findall(r'<section class="slide[^"]*"[^>]*>(.*?)</section>', page, re.S)
text = lambda s: html.unescape(re.sub(r"<[^>]+>", "", s or "")).strip()

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
blank = prs.slide_layouts[6]
for n, body in enumerate(slides, 1):
    png = SHOTS / f"slide{n:02d}.png"
    subprocess.run([BROWSER, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--window-size=1920,1080",
                    "--virtual-time-budget=6000", f"--screenshot={png}", f"{URL}?export=1#{n}"],
                   check=True, capture_output=True)
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
    h1 = text((re.search(r"<h1>(.*?)</h1>", body, re.S) or [None, ""])[1])
    notes = text((re.search(r'<aside class="notes">(.*?)</aside>', body, re.S) or [None, ""])[1])
    s.notes_slide.notes_text_frame.text = f"{h1}\n\n{notes}".strip()
    print(f"{n:2d} {h1[:70]}")
prs.save(OUT)
print(f"saved {OUT} ({len(slides)} slides)")
